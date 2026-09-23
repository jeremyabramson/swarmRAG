"""Read the Documents tab of the inventory workbook into records."""
from __future__ import annotations

from dataclasses import dataclass, field
from pathlib import Path
from typing import Iterable

from openpyxl import load_workbook

# Header text in the workbook -> attribute name
COLUMNS = {
    "Doc ID": "doc_id",
    "Technology": "technology",
    "Document title": "title",
    "Document type": "doc_type",
    "URL": "url",
    "Format": "format",
    "Access": "access",
    "Scrape method": "method",
    "Priority": "priority",
    "Top 3 rank": "top3_rank",
    "Verified": "verified",
    "What it covers / notes": "notes",
}

UNCONFIRMED_PREFIXES = ("From general knowledge", "Derived from", "Listed in the project report")


@dataclass
class DocRecord:
    doc_id: str
    technology: str
    title: str = ""
    doc_type: str = ""
    url: str = ""
    format: str = ""
    access: str = ""
    method: str = ""
    priority: int | None = None
    top3_rank: int | None = None
    verified: str = ""
    notes: str = ""
    extra: dict = field(default_factory=dict)

    @property
    def has_url(self) -> bool:
        return self.url.startswith(("http://", "https://"))

    @property
    def is_confirmed(self) -> bool:
        return self.verified.startswith(("Seen in search results", "Checked live"))

    def metadata(self) -> dict:
        return {
            "doc_id": self.doc_id,
            "technology": self.technology,
            "title": self.title,
            "doc_type": self.doc_type,
            "source_url": self.url,
            "scrape_method": self.method,
            "priority": self.priority,
            "top3_rank": self.top3_rank,
            "access": self.access,
            "url_verification": self.verified,
        }


def _to_int(value) -> int | None:
    if value in (None, ""):
        return None
    try:
        return int(value)
    except (TypeError, ValueError):
        return None


def load_records(path: str | Path, sheet: str = "Documents") -> list[DocRecord]:
    wb = load_workbook(path, read_only=True, data_only=True)
    if sheet not in wb.sheetnames:
        raise ValueError(f"Workbook has no '{sheet}' sheet; found {wb.sheetnames}")
    rows = wb[sheet].iter_rows(values_only=True)
    header = [str(h).strip() if h is not None else "" for h in next(rows)]
    missing = [h for h in ("Doc ID", "Technology", "URL", "Scrape method") if h not in header]
    if missing:
        raise ValueError(f"Documents sheet is missing columns: {missing}")
    records = []
    for row in rows:
        if not row or row[0] in (None, ""):
            continue
        values = dict(zip(header, row))
        kwargs, extra = {}, {}
        for col, val in values.items():
            attr = COLUMNS.get(col)
            if attr is None:
                extra[col] = val
            elif attr in ("priority", "top3_rank"):
                kwargs[attr] = _to_int(val)
            else:
                kwargs[attr] = "" if val is None else str(val).strip()
        records.append(DocRecord(extra=extra, **kwargs))
    wb.close()
    return records


def filter_records(
    records: Iterable[DocRecord],
    *,
    top3_only: bool = False,
    max_priority: int | None = None,
    technologies: list[str] | None = None,
    doc_ids: list[str] | None = None,
    confirmed_only: bool = False,
) -> list[DocRecord]:
    techs = {t.lower() for t in technologies} if technologies else None
    ids = set(doc_ids) if doc_ids else None
    out = []
    for r in records:
        if top3_only and not r.top3_rank:
            continue
        if max_priority is not None and (r.priority is None or r.priority > max_priority):
            continue
        if techs and not any(t in r.technology.lower() for t in techs):
            continue
        if ids and r.doc_id not in ids:
            continue
        if confirmed_only and not r.is_confirmed:
            continue
        out.append(r)
    return out
